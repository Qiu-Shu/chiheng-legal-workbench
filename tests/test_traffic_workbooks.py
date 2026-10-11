import hashlib
import io
import json
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

try:
    from traffic_workbooks import read_workbook, project_workbook
except ImportError:
    read_workbook = project_workbook = None


FIXTURES = Path(__file__).parent / "fixtures" / "traffic"


def workbook_bytes(change=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(["附件2-7"])
    sheet.append(["裁量基准"])
    sheet.append(["代码", "违法行为\n（案由）", "法律根据", "违法程度", "判断基准", "处罚基准"])
    sheet.append([7.5, "行为", "法律原文", "一般", "条件", "处罚"])
    if change:
        change(book, sheet)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def modify_zip(payload, name, transform):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload)) as source, zipfile.ZipFile(out, "w") as target:
        for entry in source.infolist():
            value = source.read(entry)
            target.writestr(entry, transform(value) if entry.filename == name else value)
    return out.getvalue()


class TrafficWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(read_workbook, "workbook reader has not been implemented")

    def read(self, change=None):
        return read_workbook(workbook_bytes(change), ".xlsx")

    def assertBlocked(self, snapshot, code):
        self.assertIn(code, [item["code"] for item in snapshot["unsupported"]])
        with self.assertRaises(ValueError):
            project_workbook(snapshot)

    def test_real_construction_preserves_hash_cells_and_display_number(self):
        payload = (FIXTURES / "construction.xlsx").read_bytes()
        snap = read_workbook(payload, ".xlsx")
        self.assertEqual(snap["sha256"], hashlib.sha256(payload).hexdigest())
        self.assertEqual(snap["unsupported"], [])
        cells = snap["sheets"][0]["cells"]
        self.assertEqual(len(cells), 1678)
        self.assertEqual(len(snap["sheets"][0]["merges"]), 349)
        for coordinate, value, display in [("A214", 7.5, "7.50"), ("A246", 7.6, "7.60"), ("A283", 7.7, "7.70"), ("A358", 7.9, "7.90")]:
            self.assertEqual(cells[coordinate]["raw_value"], value)
            self.assertEqual(cells[coordinate]["display_text"], display)
            self.assertEqual(cells[coordinate]["number_format"], "0.00_ ")
        self.assertIn("不与退还", cells["C55"]["raw_text"])

    def test_empty_real_sheet_and_its_merges_are_retained(self):
        snap = read_workbook((FIXTURES / "transport.xlsx").read_bytes(), ".xlsx")
        self.assertEqual(snap["unsupported"], [])
        self.assertEqual(len(snap["sheets"]), 2)
        self.assertEqual(snap["sheets"][1]["cells"], {})
        self.assertEqual(len(snap["sheets"][1]["merges"]), 2)

    def test_verified_government_construction_static_cells_are_complete(self):
        payload = (FIXTURES / "construction-government.xlsx").read_bytes()
        snap = read_workbook(payload, ".xlsx")
        self.assertEqual(snap["sha256"], "fcb84ac555e70bdecde03f886e1f713778b90a7af1459c8c58007c3213edf75a")
        self.assertEqual(snap["unsupported"], [])
        sheet = snap["sheets"][0]
        self.assertEqual(len(sheet["cells"]), 1678)
        self.assertEqual(len(sheet["merges"]), 349)
        self.assertIn("不予退还", sheet["cells"]["C55"]["raw_text"])
        self.assertEqual(sheet["cells"]["A214"]["display_text"], "7.50")
        text, locations = project_workbook(snap)
        self.assertIn("不予退还", text)
        self.assertEqual(len(locations), 448)
        self.assertIn("static_metadata_profile", snap["observations"])

    def test_government_metadata_support_does_not_authorize_changed_originals(self):
        original = (FIXTURES / "construction-government.xlsx").read_bytes()
        changes = [("xl/worksheets/sheet1.xml", b'filterBottomFollowUsedRange="0"', b'filterBottomFollowUsedRange="1"'),
                   ("xl/workbook.xml", b'microsoft.com:RD', b'microsoft.com:UNKNOWN'),
                   ("xl/sharedStrings.xml", "不予退还".encode(), "任意变更".encode())]
        for part, old, new in changes:
            with self.subTest(part=part):
                data = modify_zip(original, part, lambda v: v.replace(old, new))
                self.assertNotEqual(data, original)
                self.assertBlocked(read_workbook(data, ".xlsx"), "xml_structure")

    def test_xls_rich_run_prefix_does_not_strike_whole_sentence(self):
        snap = read_workbook((FIXTURES / "highway.xls").read_bytes(), ".xls")
        self.assertBlocked(snap, "biff_semantics_unverified")
        self.assertTrue(snap["observations"]["biff_records"])
        for coordinate in ("E96", "E97", "F97"):
            cell = snap["sheets"][0]["cells"][coordinate]
            runs = cell["rich_text"]
            self.assertEqual(runs[0]["text"], "\n")
            self.assertTrue(runs[0]["font"]["strike"])
            self.assertFalse(runs[1]["font"]["strike"])
            self.assertEqual("".join(run["text"] for run in runs), cell["raw_text"])

    def test_biff_font_bold_uses_weight_not_unused_flag(self):
        from xlrd.formatting import Font
        from traffic_workbook_cells import font_snapshot
        for weight, unused_flag, expected in ((700, 0, True), (700, 1, True),
                                               (400, 0, False), (400, 1, False)):
            with self.subTest(weight=weight, unused_flag=unused_flag):
                font = Font()
                font.weight, font.bold = weight, unused_flag
                self.assertEqual(font_snapshot(font, biff=True)["bold"], expected)

    def test_real_xls_header_and_footer_bold_weight_is_preserved(self):
        snap = read_workbook((FIXTURES / "highway.xls").read_bytes(), ".xls")
        cells = snap["sheets"][0]["cells"]
        for coordinate in ("A3", "B3", "C3", "D3", "E3", "F3"):
            self.assertTrue(cells[coordinate]["font"]["bold"], coordinate)
        self.assertTrue(cells["A172"]["rich_text"][0]["font"]["bold"])
        self.assertFalse(cells["A172"]["rich_text"][1]["font"]["bold"])
        self.assertBlocked(snap, "biff_semantics_unverified")

    def test_projection_context_coordinates_and_ordinary_blank(self):
        def change(book, sheet):
            sheet.append([7.5, None, None, "严重", "条件二", "处罚二"])
            sheet.merge_cells("C4:C5")
        snap = self.read(change)
        text, locations = project_workbook(snap)
        row = next(loc for loc in locations if loc["row"] == 5)
        self.assertEqual(len(row["cells"]), 6)
        self.assertEqual(row["cells"][1]["raw_text"], "")
        self.assertIsNone(row["cells"][1]["merge_anchor"])
        self.assertEqual(row["cells"][2]["source_coordinate"], "C4")
        self.assertEqual(row["cells"][2]["raw_text"], "法律原文")
        self.assertIn("C5", text[row["start"]:row["end"]])
        self.assertIn("C4", text[row["start"]:row["end"]])
        self.assertNotEqual(row["id"], next(loc for loc in locations if loc["row"] == 4)["id"])
        for location in locations:
            for entry in location["cells"]:
                original = snap["sheets"][0]["cells"].get(entry["source_coordinate"])
                self.assertEqual(entry["raw_text"], original["raw_text"] if original else "")

    def test_formulas_and_errors_are_preserved_but_blocked(self):
        for value, code in [("=1+1", "formula"), ("#DIV/0!", "error_value")]:
            with self.subTest(value=value):
                snap = self.read(lambda b, s: setattr(s["E4"], "value", value))
                self.assertEqual(snap["sheets"][0]["cells"]["E4"]["raw_value"], value)
                self.assertBlocked(snap, code)

    def test_hidden_sheet_row_column_are_blocked(self):
        changes = [lambda b, s: setattr(b.create_sheet("hidden"), "sheet_state", "hidden"),
                   lambda b, s: setattr(s.row_dimensions[4], "hidden", True),
                   lambda b, s: setattr(s.column_dimensions["E"], "hidden", True)]
        for change in changes:
            with self.subTest(change=change):
                self.assertBlocked(self.read(change), "hidden_content")

    def test_header_footer_content_is_observed_and_blocked(self):
        snap = self.read(lambda b, s: setattr(s.oddHeader.center, "text", "附加法律文字"))
        self.assertIn("附加法律文字", snap["sheets"][0]["observations"]["header_footer"])
        self.assertBlocked(snap, "header_footer")

    def test_unknown_number_format_is_blocked(self):
        self.assertBlocked(self.read(lambda b, s: setattr(s["A4"], "number_format", "0.0%")), "number_format")

    def test_real_highway_numeric_only_format_does_not_change_text(self):
        snap = read_workbook((FIXTURES / "highway.xls").read_bytes(), ".xls")
        selected = [cell for cell in snap["sheets"][0]["cells"].values()
                    if cell["number_format"] == r'0\.0,"万"']
        self.assertEqual(len(selected), 168)
        for cell in selected:
            self.assertEqual(cell["value_type"], "s")
            self.assertEqual(cell["display_text"], cell["raw_text"])
        self.assertNotIn("number_format", [x["code"] for x in snap["unsupported"]])
        self.assertBlocked(snap, "biff_semantics_unverified")

    def test_verified_numeric_only_format_is_supported_for_text_not_numbers(self):
        for value in ("处1000元以下的罚款", "1000", "7.50"):
            with self.subTest(value=value):
                def change(book, sheet):
                    sheet["E4"] = value
                    sheet["E4"].number_format = r'0\.0,"万"'
                snap = self.read(change)
                self.assertEqual(snap["unsupported"], [])
                cell = snap["sheets"][0]["cells"]["E4"]
                self.assertEqual(cell["raw_value"], value)
                self.assertEqual(cell["display_text"], value)
                text, locations = project_workbook(snap)
                self.assertIn(value, text)
        self.assertBlocked(self.read(lambda b, s: setattr(s["A4"], "number_format", r'0\.0,"万"')), "number_format")

    def test_numeric_only_text_format_does_not_allow_unverified_variants(self):
        for fmt in (r'0\.0,"万";;;', r'0\.0,"万";;;"前缀"@',
                    r'[White]0\.0,"万"', r'0\.0,"万"@'):
            with self.subTest(number_format=fmt):
                self.assertBlocked(self.read(lambda b, s: setattr(s["E4"], "number_format", fmt)), "number_format")

    def test_rich_text_semantics_preserved_and_blocked(self):
        for font in (InlineFont(strike=True), InlineFont(vertAlign="superscript")):
            value = CellRichText("普通", TextBlock(font, "特殊"))
            snap = self.read(lambda b, s: setattr(s["E4"], "value", value))
            self.assertEqual(snap["sheets"][0]["cells"]["E4"]["raw_text"], "普通特殊")
            self.assertEqual(snap["sheets"][0]["cells"]["E4"]["rich_text"][1]["text"], "特殊")
            self.assertBlocked(snap, "font_semantics")

    def test_plain_rich_text_is_retained(self):
        snap = self.read(lambda b, s: setattr(s["E4"], "value", CellRichText("普通", TextBlock(InlineFont(b=True), "加粗"))))
        self.assertEqual(snap["unsupported"], [])
        self.assertTrue(snap["sheets"][0]["cells"]["E4"]["rich_text"][1]["font"]["bold"])

    def test_merge_cannot_silently_discard_nonempty_subcell(self):
        data = modify_zip(workbook_bytes(), "xl/worksheets/sheet1.xml", lambda v: v.replace(b"</sheetData>", b'</sheetData><mergeCells count="1"><mergeCell ref="D4:E4"/></mergeCells>'))
        snap = read_workbook(data, ".xlsx")
        self.assertIn("E4", snap["sheets"][0]["cells"])
        self.assertBlocked(snap, "merge_nonempty_child")

    def test_overlap_merges_blocked(self):
        data = modify_zip(workbook_bytes(), "xl/worksheets/sheet1.xml", lambda v: v.replace(b"</sheetData>", b'</sheetData><mergeCells count="2"><mergeCell ref="A5:B6"/><mergeCell ref="B6:C7"/></mergeCells>'))
        self.assertBlocked(read_workbook(data, ".xlsx"), "merge_overlap")

    def test_external_link_blocked(self):
        self.assertBlocked(self.read(lambda b, s: setattr(s["E4"], "hyperlink", "https://example.com")), "external_relationship")

    def test_drawing_macro_unknown_package_are_blocked(self):
        for name in ("xl/drawings/drawing1.xml", "xl/vbaProject.bin", "xl/unknown.xml"):
            out = io.BytesIO(workbook_bytes())
            with zipfile.ZipFile(out, "a") as archive:
                archive.writestr(name, "<unknown/>")
            self.assertBlocked(read_workbook(out.getvalue(), ".xlsx"), "package_part")

    def test_unknown_worksheet_extension_blocked(self):
        data = modify_zip(workbook_bytes(), "xl/worksheets/sheet1.xml", lambda v: v.replace(b"</worksheet>", b'<extLst><ext uri="unknown"/></extLst></worksheet>'))
        self.assertBlocked(read_workbook(data, ".xlsx"), "xml_structure")

    def test_header_mutation_and_content_beyond_six_columns_blocked(self):
        self.assertBlocked(self.read(lambda b, s: setattr(s["C3"], "value", "新表头")), "header_structure")
        self.assertBlocked(self.read(lambda b, s: setattr(s["G4"], "value", "第七列原文")), "table_columns")

    def test_corrupt_and_mismatched_files_fail_closed(self):
        for payload, suffix in [(b"garbage", ".xlsx"), (b"PK\x03\x04", ".xlsx"), (workbook_bytes(), ".xls"), (workbook_bytes(), ".xlsm")]:
            with self.subTest(suffix=suffix), self.assertRaises(ValueError):
                read_workbook(payload, suffix)

    def test_bounded_input_and_worksheet_dimensions(self):
        with self.assertRaises(ValueError):
            read_workbook(b"x" * (10 * 1024 * 1024 + 1), ".xlsx")
        data = modify_zip(workbook_bytes(), "xl/worksheets/sheet1.xml", lambda v: v.replace(b'A1:F4', b'A1:XFD1048576'))
        with self.assertRaises(ValueError):
            read_workbook(data, ".xlsx")

    def test_text_hidden_by_custom_format_is_blocked(self):
        self.assertBlocked(self.read(lambda b, s: setattr(s["E4"], "number_format", ";;;")), "number_format")

    def test_unknown_attributes_and_misplaced_known_elements_are_blocked(self):
        for old, new in [(b'<sheetData>', b'<sheetData future="1">'),
                         (b'</sheetData>', '</sheetData><si><t>遗漏原文</t></si>'.encode())]:
            data = modify_zip(workbook_bytes(), "xl/worksheets/sheet1.xml", lambda v: v.replace(old, new))
            self.assertBlocked(read_workbook(data, ".xlsx"), "xml_structure")

    def test_zero_height_default_and_plain_font_semantics_are_blocked(self):
        self.assertBlocked(self.read(lambda b, s: setattr(s.sheet_format, "zeroHeight", True)), "hidden_content")
        from openpyxl.styles import Font
        self.assertBlocked(self.read(lambda b, s: setattr(s["E4"], "font", Font(strike=True))), "font_semantics")

    def test_duplicate_coordinates_and_excessive_merges_raise(self):
        for extra in [b'<row r="4"><c r="E4" t="n"><v>1</v></c></row>',
                      b'<row r="10001"><c r="A10001" t="n"><v>1</v></c></row>']:
            data = modify_zip(workbook_bytes(), "xl/worksheets/sheet1.xml", lambda v: v.replace(b'</sheetData>', extra + b'</sheetData>'))
            with self.assertRaises(ValueError):
                read_workbook(data, ".xlsx")

    def test_package_relationship_type_and_shared_strings_extension_blocked(self):
        data = modify_zip(workbook_bytes(), "xl/_rels/workbook.xml.rels", lambda v: v.replace(b'/worksheet"', b'/unknown"'))
        self.assertBlocked(read_workbook(data, ".xlsx"), "package_relationship")

    def test_projection_offsets_are_exact_and_reproducible(self):
        snap = self.read()
        first = project_workbook(snap)
        self.assertEqual(first, project_workbook(snap))
        text, locations = first
        self.assertEqual(len({loc["id"] for loc in locations}), len(locations))
        self.assertEqual("".join(text[loc["start"]:loc["end"]] for loc in locations), text)

    def test_zero_size_dimensions_blocked(self):
        for field, attribute in [("row", "ht"), ("col", "width")]:
            def change(book, sheet):
                if field == "row":
                    sheet.row_dimensions[4].height = 0
                else:
                    sheet.column_dimensions["E"].width = 0.001
            payload = workbook_bytes(change)
            if field == "row":
                payload = modify_zip(payload, "xl/worksheets/sheet1.xml", lambda v: v.replace(b'customHeight="1"', b'customHeight="1" ht="0"'))
            else:
                payload = modify_zip(payload, "xl/worksheets/sheet1.xml", lambda v: v.replace(b'width="0.001"', b'width="0"'))
            self.assertBlocked(read_workbook(payload, ".xlsx"), "hidden_content")

    def test_xml_entities_and_duplicate_package_parts_are_rejected(self):
        data = modify_zip(workbook_bytes(), "xl/worksheets/sheet1.xml", lambda v: b'<!DOCTYPE worksheet [<!ENTITY hidden "value">]>' + v)
        with self.assertRaises(ValueError):
            read_workbook(data, ".xlsx")
        data = modify_zip(workbook_bytes(), "xl/worksheets/sheet1.xml", lambda v: ('<!DOCTYPE worksheet [<!ENTITY hidden "value">]>' + v.decode()).encode('utf-16'))
        with self.assertRaises(ValueError):
            read_workbook(data, ".xlsx")

    def test_content_type_mismatch_and_unknown_relationship_elements_blocked(self):
        data = modify_zip(workbook_bytes(), "[Content_Types].xml", lambda v: v.replace(b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml', b'application/vnd.ms-excel.sheet.macroEnabled.main+xml'))
        self.assertBlocked(read_workbook(data, ".xlsx"), "package_content_type")
        data = modify_zip(workbook_bytes(), "xl/_rels/workbook.xml.rels", lambda v: v.replace(b'</Relationships>', b'<unknown/></Relationships>'))
        self.assertBlocked(read_workbook(data, ".xlsx"), "package_relationship")

    def test_projection_empty_sheet_and_full_width_notes_remain_visible(self):
        def change(book, sheet):
            sheet["A5"] = "备注原文"
            sheet.merge_cells("A5:F5")
            empty = book.create_sheet("Empty")
            empty.merge_cells("A1:F1")
        text, locations = project_workbook(self.read(change))
        note = next(loc for loc in locations if loc["row"] == 5)
        self.assertEqual(text[note["start"]:note["end"]].count("备注原文"), 1)
        self.assertNotIn("处罚基准", text[note["start"]:note["end"]])
        self.assertTrue(any(loc["sheet"] == "Empty" for loc in locations))
        self.assertIn("A1:F1", text)
        regular = next(loc for loc in locations if loc["row"] == 4)
        self.assertIn("法律根据", text[regular["start"]:regular["end"]])

    def test_inherited_row_and_column_styles_fail_closed(self):
        from openpyxl.styles import Font
        for change in [lambda b, s: setattr(s.row_dimensions[4], "number_format", "0.00"),
                       lambda b, s: setattr(s.column_dimensions["A"], "number_format", "0.00"),
                       lambda b, s: setattr(s.row_dimensions[4], "font", Font(strike=True)),
                       lambda b, s: setattr(s.column_dimensions["E"], "number_format", ";;;")]:
            self.assertBlocked(self.read(change), "style_inheritance")

    def test_duplicate_values_nonempty_xml_text_and_relationship_roots_blocked(self):
        changes = [
            ("xl/worksheets/sheet1.xml", lambda v: v.replace(b'<v>7.5</v>', b'<v>7.5</v><v>999</v>')),
            ("xl/worksheets/sheet1.xml", lambda v: v.replace(b'<sheetData>', b'<sheetData>LOST LEGAL TEXT')),
            ("xl/_rels/workbook.xml.rels", lambda v: v.replace(b'Relationships', b'UNKNOWN')),
        ]
        for part, transform in changes:
            snapshot = read_workbook(modify_zip(workbook_bytes(), part, transform), ".xlsx")
            self.assertTrue(snapshot["unsupported"])
            with self.assertRaises(ValueError):
                project_workbook(snapshot)

    def test_package_xml_unknown_children_tags_text_and_root_attributes_blocked(self):
        changes = [
            ("xl/_rels/workbook.xml.rels", lambda v: v.replace(b'/>', b'><unknown>LOST LEGAL TEXT</unknown></Relationship>', 1)),
            ("[Content_Types].xml", lambda v: v.replace(b'<Override ', b'<Unknown ', 1)),
            ("[Content_Types].xml", lambda v: v.replace(b'<Types ', b'<Types future="1" ', 1)),
            ("xl/_rels/workbook.xml.rels", lambda v: v.replace(b'<Relationships ', b'<Relationships future="1" ', 1)),
            ("[Content_Types].xml", lambda v: v.replace(b'/><', b'/>LOST LEGAL TEXT<', 1)),
        ]
        for part, transform in changes:
            with self.subTest(part=part, transform=transform):
                snapshot = read_workbook(modify_zip(workbook_bytes(), part, transform), ".xlsx")
                self.assertTrue(snapshot["unsupported"])
                with self.assertRaises(ValueError):
                    project_workbook(snapshot)

    def test_conflicting_cell_representations_are_preserved_and_blocked(self):
        original = b'<c r="A4" t="n"><v>7.5</v></c>'
        for replacement in [
            b'<c r="A4" t="n"><v>7.5</v><is><t>LOST LEGAL TEXT</t></is></c>',
            b'<c r="A4" t="inlineStr"><v>7.5</v><is><t>LOST LEGAL TEXT</t></is></c>',
            b'<c r="A4" t="n"><f>1+1</f><is><t>LOST LEGAL TEXT</t></is></c>',
        ]:
            with self.subTest(replacement=replacement):
                payload = modify_zip(workbook_bytes(), "xl/worksheets/sheet1.xml", lambda v: v.replace(original, replacement))
                snapshot = read_workbook(payload, ".xlsx")
                self.assertBlocked(snapshot, "cell_representation")
                self.assertIn("LOST LEGAL TEXT", snapshot["sheets"][0]["cells"]["A4"]["source_xml"])

    def test_noncanonical_boolean_values_are_not_coerced_to_false(self):
        original = b'<c r="A4" t="n"><v>7.5</v></c>'
        for value in ("2", "garbage", "true"):
            with self.subTest(value=value):
                replacement = f'<c r="A4" t="b"><v>{value}</v></c>'.encode()
                payload = modify_zip(workbook_bytes(), "xl/worksheets/sheet1.xml", lambda v: v.replace(original, replacement))
                snapshot = read_workbook(payload, ".xlsx")
                self.assertBlocked(snapshot, "cell_representation")
                self.assertEqual(snapshot["sheets"][0]["cells"]["A4"]["raw_value"], value)

    def test_real_workbook_snapshots_round_trip_strict_json(self):
        for name in ("construction.xlsx", "transport.xlsx", "highway.xls"):
            with self.subTest(name=name):
                path = FIXTURES / name
                snapshot = read_workbook(path.read_bytes(), path.suffix)
                encoded = json.dumps(snapshot, ensure_ascii=False, allow_nan=False)
                self.assertEqual(json.loads(encoded), snapshot)

    def test_rich_text_font_name_is_serializable_and_preserved(self):
        value = CellRichText(TextBlock(InlineFont(rFont="宋体", b=True), "字体原文"))
        snapshot = self.read(lambda b, s: setattr(s["E4"], "value", value))
        self.assertEqual(snapshot["sheets"][0]["cells"]["E4"]["rich_text"][0]["font"]["name"], "宋体")
        json.dumps(snapshot, ensure_ascii=False, allow_nan=False)


if __name__ == "__main__":
    unittest.main()
