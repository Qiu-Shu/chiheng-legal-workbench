"""Read-only, bounded extraction for the verified six-column traffic tables.

This is deliberately not a general Excel renderer. Unknown structures are
recorded and prevent projection. XLS cached values remain unverified.
"""

import hashlib
import io
import math
import posixpath
import re
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
import xlrd
from openpyxl.styles import Font
from openpyxl.styles.numbers import BUILTIN_FORMATS
from openpyxl.styles.stylesheet import Stylesheet
from openpyxl.cell.text import InlineFont
from openpyxl.utils.cell import get_column_letter, range_boundaries, coordinate_to_tuple

from traffic_workbook_cells import cell_snapshot, font_snapshot

MAX_INPUT = 10 * 1024 * 1024
MAX_EXPANDED = 50 * 1024 * 1024
MAX_ROWS, MAX_COLUMNS, MAX_CELLS = 10000, 64, 200000
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
HEADERS = ("代码", "违法行为（案由）", "法律根据", "违法程度", "判断基准", "处罚基准")

# Metadata/layout elements accepted in the inspected sources. Unknown attributes
# also fail closed, because parsers otherwise silently ignore their semantics.
ATTRS = {
    "worksheet": "", "sheetPr": "codeName", "outlinePr": "summaryBelow summaryRight showOutlineSymbols applyStyles",
    "pageSetUpPr": "autoPageBreaks fitToPage", "dimension": "ref", "sheetViews": "",
    "sheetView": "tabSelected topLeftCell zoomScaleNormal workbookViewId view zoomScaleSheetLayoutView zoomScale showGridLines showRowColHeaders rightToLeft",
    "selection": "pane activeCell activeCellId sqref", "pane": "xSplit ySplit topLeftCell activePane state",
    "sheetFormatPr": "outlineLevelCol outlineLevelRow defaultRowHeight defaultColWidth baseColWidth customHeight zeroHeight thickTop thickBottom",
    "cols": "", "col": "customWidth min width max style hidden bestFit outlineLevel collapsed",
    "sheetData": "", "row": "customHeight s r customFormat ht spans hidden outlineLevel collapsed thickTop thickBot",
    "c": "s r t", "v": "", "f": "t ref si ca aca bx dt2D dtr del1 del2 r1 r2",
    "is": "", "t": "{http://www.w3.org/XML/1998/namespace}space", "r": "", "rPr": "",
    "sheetProtection": "selectUnlockedCells selectLockedCells sheet objects scenarios password formatCells formatColumns formatRows insertColumns insertRows insertHyperlinks deleteColumns deleteRows sort autoFilter pivotTables",
    "autoFilter": "ref", "mergeCells": "count", "mergeCell": "ref",
    "printOptions": "horizontalCentered verticalCentered headings gridLines gridLinesSet",
    "pageMargins": "bottom left footer top header right", "pageSetup": "scale horizontalDpi verticalDpi orientation paperSize fitToHeight fitToWidth firstPageNumber useFirstPageNumber blackAndWhite draft cellComments errors copies",
    "headerFooter": "scaleWithDoc alignWithMargins differentOddEven differentFirst", "oddHeader": "", "oddFooter": "", "evenHeader": "", "evenFooter": "", "firstHeader": "", "firstFooter": "",
    "rowBreaks": "manualBreakCount count", "colBreaks": "manualBreakCount count", "brk": "id min max man pt",
    "ignoredErrors": "", "ignoredError": "numberStoredAsText sqref evalError twoDigitTextYear formula formulaRange unlockedFormula emptyCellReference listDataValidation calculatedColumn",
    "hyperlinks": "", "hyperlink": "ref location display tooltip {" + REL + "}id",
    "extLst": "", "ext": "uri", "sst": "uniqueCount count", "si": "",
    "workbook": "", "fileVersion": "lowestEdited rupBuild appName lastEdited", "workbookPr": "date1904 defaultThemeVersion codeName",
    "bookViews": "", "workbookView": "windowWidth windowHeight xWindow yWindow tabRatio visibility minimized showHorizontalScroll showVerticalScroll showSheetTabs firstSheet activeTab autoFilterDateGrouping",
    "sheets": "", "sheet": "sheetId name state {" + REL + "}id", "definedNames": "", "definedName": "name localSheetId hidden",
    "calcPr": "calcId fullCalcOnLoad", "workbookProtection": "", "styleSheet": "", "numFmts": "count", "numFmt": "formatCode numFmtId",
    "fonts": "count", "font": "", "fills": "count", "fill": "", "patternFill": "patternType",
    "fgColor": "theme rgb tint indexed auto", "bgColor": "theme rgb tint indexed auto", "color": "theme rgb indexed tint auto",
    "borders": "count", "border": "diagonalUp diagonalDown outline", "left": "style", "right": "style", "top": "style", "bottom": "style", "diagonal": "style", "start": "style", "end": "style",
    "cellStyleXfs": "count", "cellXfs": "count", "xf": "applyFill borderId xfId applyFont numFmtId applyProtection fillId applyNumberFormat applyAlignment applyBorder fontId pivotButton quotePrefix",
    "alignment": "vertical horizontal wrapText textRotation shrinkToFit indent relativeIndent justifyLastLine readingOrder", "protection": "locked hidden",
    "cellStyles": "count", "cellStyle": "builtinId xfId name hidden customBuiltin", "tableStyles": "defaultPivotStyle defaultTableStyle count", "colors": "", "mruColors": "", "indexedColors": "", "rgbColor": "rgb",
}
for _tag in ("sz", "name", "charset", "b", "i", "u", "scheme", "rFont", "family", "strike", "vertAlign", "outline", "shadow", "condense", "extend"):
    ATTRS[_tag] = "val"

CHILDREN = {
    "worksheet": "sheetPr dimension sheetViews sheetFormatPr cols sheetData sheetProtection autoFilter mergeCells printOptions pageMargins pageSetup headerFooter rowBreaks colBreaks ignoredErrors hyperlinks extLst",
    "sheetPr": "outlinePr pageSetUpPr", "sheetViews": "sheetView", "sheetView": "selection pane",
    "cols": "col", "sheetData": "row", "row": "c", "c": "v f is", "is": "t r",
    "r": "rPr t", "si": "t r", "sst": "si", "mergeCells": "mergeCell", "autoFilter": "extLst",
    "headerFooter": "oddHeader oddFooter evenHeader evenFooter firstHeader firstFooter",
    "rowBreaks": "brk", "colBreaks": "brk", "ignoredErrors": "ignoredError", "hyperlinks": "hyperlink",
    "workbook": "fileVersion workbookPr workbookProtection bookViews sheets definedNames calcPr",
    "bookViews": "workbookView", "sheets": "sheet", "definedNames": "definedName",
    "styleSheet": "numFmts fonts fills borders cellStyleXfs cellXfs cellStyles tableStyles colors extLst",
    "numFmts": "numFmt", "fonts": "font", "fills": "fill", "fill": "patternFill",
    "patternFill": "fgColor bgColor", "borders": "border", "border": "left right top bottom diagonal start end",
    "cellStyleXfs": "xf", "cellXfs": "xf", "xf": "alignment protection", "cellStyles": "cellStyle",
    "colors": "mruColors indexedColors", "mruColors": "color", "indexedColors": "rgbColor",
    "extLst": "ext", "ext": "slicerStyles",
}
for _tag in ("font", "rPr"):
    CHILDREN[_tag] = "sz name charset b i u scheme rFont family strike vertAlign outline shadow condense extend color"
for _tag in ("left", "right", "top", "bottom", "diagonal", "start", "end"):
    CHILDREN[_tag] = "color"


def _flag(snapshot, code, **context):
    item = {"code": code, **context}
    if item not in snapshot["unsupported"]:
        snapshot["unsupported"].append(item)


def _xml(data):
    if b"\x00" in data or b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
        raise ValueError("DTD/entity declarations are unsupported")
    return ET.fromstring(data)


def _verified_static_metadata(root, snapshot, part):
    """One independently compared original, never a general extension whitelist.

    All cells, rich runs, fonts and merges in this exact source were compared
    with openpyxl. Its full XML remains in observations. The WPS flag is only
    observed in an inactive filter; no filtering or calculation is performed.
    A changed byte has a different SHA and receives the ordinary strict audit.
    """
    digest = "fcb84ac555e70bdecde03f886e1f713778b90a7af1459c8c58007c3213edf75a"
    if snapshot["sha256"] != digest:
        return set(), set()
    mc = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
    nodes, attributes = set(), set()
    expected = {"xl/styles.xml": "xr9", "xl/workbook.xml": "x15"}.get(part)
    if expected and root.get(mc + "Ignorable") == expected:
        attributes.add(mc + "Ignorable")
    if part == "xl/workbook.xml":
        for child in root:
            if child.tag in (mc + "AlternateContent", f"{{{NS}}}extLst"):
                nodes.update(child.iter())
    elif part == "xl/worksheets/sheet1.xml":
        element = root.find(f"{{{NS}}}autoFilter")
        wps = "{http://www.wps.cn/officeDocument/2017/etCustomData}filterBottomFollowUsedRange"
        if (element is not None and element.attrib == {"ref": "A3:F448", wps: "0"}
                and len(element) == 1 and element[0].tag == f"{{{NS}}}extLst"
                and not element[0].attrib and not len(element[0])):
            nodes.update(element.iter())
    if nodes or attributes:
        profile = snapshot["observations"].setdefault("static_metadata_profile", {
            "id": "government-construction-static-2026-10-03", "original_sha256": digest,
            "scope": "静态全格提取；私有筛选UI行为未验收；字节变化须重新核验", "parts": {}})
        profile["parts"][part] = {"root_attributes": {key: root.get(key) for key in attributes},
                                  "retained_xml": [ET.tostring(child, encoding="unicode")
                                                   for child in root if child in nodes]}
    return nodes, attributes


def _audit_xml(root, snapshot, part):
    metadata_nodes, root_attributes = _verified_static_metadata(root, snapshot, part)
    for element in root.iter():
        if element in metadata_nodes:
            continue
        name = element.tag.removeprefix("{" + NS + "}")
        # Office's default slicer style is metadata, not a slicer or its content.
        if (part == "xl/styles.xml" and element.tag == "{http://schemas.microsoft.com/office/spreadsheetml/2009/9/main}slicerStyles"
                and not list(element) and set(element.attrib) == {"defaultSlicerStyle"}):
            continue
        permitted = set(ATTRS.get(name, "").split()) | (root_attributes if element is root else set())
        if name not in ATTRS or set(element.attrib) - permitted:
            _flag(snapshot, "xml_structure", part=part, element=element.tag, attributes=dict(element.attrib))
        text_nodes = {"t", "v", "f", "definedName", "oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter"}
        if (element.text and element.text.strip() and name not in text_nodes) or (element.tail and element.tail.strip()):
            _flag(snapshot, "xml_structure", part=part, element=element.tag, detail="unexpected text")
        repeatable = {"sst": {"si"}, "si": {"r"}, "is": {"r"}, "sheetData": {"row"}, "row": {"c"},
                      "cols": {"col"}, "mergeCells": {"mergeCell"}, "sheetViews": {"sheetView"},
                      "sheetView": {"selection"}, "sheets": {"sheet"}, "bookViews": {"workbookView"},
                      "definedNames": {"definedName"}, "numFmts": {"numFmt"}, "fonts": {"font"},
                      "fills": {"fill"}, "borders": {"border"}, "cellXfs": {"xf"}, "cellStyleXfs": {"xf"},
                      "cellStyles": {"cellStyle"}, "mruColors": {"color"}, "indexedColors": {"rgbColor"},
                      "rowBreaks": {"brk"}, "colBreaks": {"brk"}, "ignoredErrors": {"ignoredError"},
                      "hyperlinks": {"hyperlink"}, "extLst": {"ext"}}
        seen = set()
        for child in element:
            if child in metadata_nodes:
                continue
            child_name = child.tag.rsplit("}", 1)[-1]
            if child_name not in CHILDREN.get(name, "").split():
                _flag(snapshot, "xml_structure", part=part, element=child.tag, parent=element.tag)
            if child_name in seen and child_name not in repeatable.get(name, set()):
                _flag(snapshot, "xml_structure", part=part, element=child.tag, detail="duplicate singleton")
            seen.add(child_name)
        if name == "ext" and not (part == "xl/styles.xml" and element.get("uri") == "{EB79DEF2-80B8-43e5-95BD-54CBDDF9020C}"):
            _flag(snapshot, "xml_structure", part=part, element=element.tag)


def _bounds(reference):
    bounds = range_boundaries(reference)
    if (any(x is None or x < 1 for x in bounds) or bounds[0] > bounds[2] or bounds[1] > bounds[3]
            or bounds[2] > MAX_COLUMNS or bounds[3] > MAX_ROWS
            or (bounds[2] - bounds[0] + 1) * (bounds[3] - bounds[1] + 1) > MAX_CELLS):
        raise ValueError("Worksheet range exceeds supported limits")
    return bounds


def _audit_package_xml(root, snapshot, part):
    """The two package manifests contain only a root and attribute-only leaves."""
    if root.attrib or (root.text and root.text.strip()):
        _flag(snapshot, "xml_structure", part=part, detail="package root attributes or text")
    if part.endswith(".rels"):
        allowed = {f"{{{PKG_REL}}}Relationship": ({"Id", "Type", "Target"}, {"TargetMode"})}
    else:
        prefix = "{http://schemas.openxmlformats.org/package/2006/content-types}"
        allowed = {prefix + "Default": ({"Extension", "ContentType"}, set()),
                   prefix + "Override": ({"PartName", "ContentType"}, set())}
    for leaf in root:
        required, optional = allowed.get(leaf.tag, (set(), set()))
        if (leaf.tag not in allowed or list(leaf) or (leaf.text and leaf.text.strip())
                or (leaf.tail and leaf.tail.strip()) or set(leaf.attrib) - required - optional
                or any(not leaf.get(key) for key in required)):
            _flag(snapshot, "xml_structure", part=part, element=leaf.tag, detail="unsupported package leaf")


def _check_cell(snapshot, sheet, cell):
    context = {"sheet": sheet["name"], "coordinate": cell["coordinate"]}
    if cell["value_type"] in ("f", "e", "d"):
        _flag(snapshot, {"f": "formula", "e": "error_value", "d": "date_value"}[cell["value_type"]], **context)
    # This verified single numeric section has no text section; numbers stay unsupported.
    text_with_numeric_format = (cell["value_type"] in ("s", "inlineStr")
                                and cell["number_format"] == r'0\.0,"万"')
    if cell["display_text"] is None or (cell["number_format"] not in ("General", "@") and not re.fullmatch(r"0(?:\.(0+))?(?:_ )?", cell["number_format"]) and not text_with_numeric_format):
        _flag(snapshot, "number_format", number_format=cell["number_format"], **context)
    for run in cell["rich_text"] or [{"text": cell["raw_text"], "font": cell["font"]}]:
        if run["text"].strip() and (run["font"]["strike"] or run["font"]["vertAlign"]):
            _flag(snapshot, "font_semantics", **context)


def _finish_sheet(snapshot, sheet):
    merged_cells = {}
    total_area = 0
    for reference in sheet["merges"]:
        left, top, right, bottom = _bounds(reference)
        total_area += (right - left + 1) * (bottom - top + 1)
        if total_area > MAX_CELLS:
            raise ValueError("Merged areas exceed supported limits")
        anchor = f"{get_column_letter(left)}{top}"
        for row in range(top, bottom + 1):
            for col in range(left, right + 1):
                coordinate = f"{get_column_letter(col)}{row}"
                if coordinate in merged_cells:
                    _flag(snapshot, "merge_overlap", sheet=sheet["name"], reference=reference)
                merged_cells[coordinate] = anchor
                if coordinate in sheet["cells"] and coordinate != anchor:
                    _flag(snapshot, "merge_nonempty_child", sheet=sheet["name"], coordinate=coordinate, anchor=anchor)
    if sheet["visibility"] not in ("visible", 0) or sheet["hidden_rows"] or sheet["hidden_columns"] or sheet.get("default_row_hidden"):
        _flag(snapshot, "hidden_content", sheet=sheet["name"])
    for cell in sheet["cells"].values():
        _check_cell(snapshot, sheet, cell)
        if coordinate_to_tuple(cell["coordinate"])[1] > 6:
            _flag(snapshot, "table_columns", sheet=sheet["name"], coordinate=cell["coordinate"])
    if sheet["cells"]:
        headers = [re.sub(r"\s+", "", sheet["cells"].get(f"{c}3", {}).get("raw_text", "")) for c in "ABCDEF"]
        if tuple(headers) != HEADERS:
            _flag(snapshot, "header_structure", sheet=sheet["name"], actual=headers)


def _rich_text(element, base_font):
    runs = []
    for child in element:
        if child.tag == f"{{{NS}}}t":
            runs.append({"text": child.text or "", "font": base_font})
        elif child.tag == f"{{{NS}}}r":
            props = child.find(f"{{{NS}}}rPr")
            font = font_snapshot(InlineFont.from_tree(props)) if props is not None else base_font
            runs.append({"text": child.findtext(f"{{{NS}}}t", ""), "font": font})
    return "".join(run["text"] for run in runs), runs


def _read_xlsx(payload, snapshot):
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        entries = archive.infolist()
        if len(entries) > 1000 or sum(e.file_size for e in entries) > MAX_EXPANDED:
            raise ValueError("Expanded workbook exceeds supported limits")
        names = [entry.filename for entry in entries if not entry.is_dir()]
        if len(names) != len(set(names)):
            raise ValueError("Duplicate package parts")
        parts = {name: archive.read(name) for name in names}
    allowed = r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core|custom)\.xml|xl/(?:workbook|styles|sharedStrings)\.xml|xl/_rels/workbook\.xml\.rels|xl/theme/theme1\.xml|xl/worksheets/sheet[0-9]+\.xml|xl/worksheets/_rels/sheet[0-9]+\.xml\.rels)"
    snapshot["observations"]["package_parts"] = names
    roots = {}
    for name, data in parts.items():
        if not re.fullmatch(allowed, name):
            _flag(snapshot, "package_part", part=name)
        if name.endswith((".xml", ".rels")):
            roots[name] = _xml(data)
            expected_root = None
            if name.endswith(".rels"):
                expected_root = f"{{{PKG_REL}}}Relationships"
            elif name == "[Content_Types].xml":
                expected_root = "{http://schemas.openxmlformats.org/package/2006/content-types}Types"
            elif name in {"xl/workbook.xml", "xl/styles.xml", "xl/sharedStrings.xml"}:
                expected_root = f"{{{NS}}}" + {"xl/workbook.xml": "workbook", "xl/styles.xml": "styleSheet", "xl/sharedStrings.xml": "sst"}[name]
            elif re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml", name):
                expected_root = f"{{{NS}}}worksheet"
            if expected_root and roots[name].tag != expected_root:
                _flag(snapshot, "xml_structure", part=name, detail="unexpected XML root", element=roots[name].tag)
            if name.endswith(".rels") or name == "[Content_Types].xml":
                _audit_package_xml(roots[name], snapshot, name)
                snapshot["observations"].setdefault("package_xml", {})[name] = data.decode("utf-8-sig")
        if name.endswith(".rels"):
            for relationship in roots[name]:
                if relationship.tag != f"{{{PKG_REL}}}Relationship" or set(relationship.attrib) - {"Id", "Type", "Target", "TargetMode"}:
                    _flag(snapshot, "package_relationship", part=name, relationship=dict(relationship.attrib))
                if relationship.get("TargetMode") == "External":
                    _flag(snapshot, "external_relationship", part=name, relationship=dict(relationship.attrib))
                kind = relationship.get("Type", "").rsplit("/", 1)[-1]
                if kind not in {"worksheet", "styles", "sharedStrings", "theme", "officeDocument", "core-properties", "extended-properties", "custom-properties"}:
                    _flag(snapshot, "package_relationship", part=name, relationship=dict(relationship.attrib))
        if name.startswith("xl/") and name.endswith(".xml") and not name.startswith("xl/theme/"):
            _audit_xml(roots[name], snapshot, name)
    accepted_content_types = {"application/xml", "application/vnd.openxmlformats-package.relationships+xml",
                              "application/vnd.openxmlformats-package.core-properties+xml"}
    accepted_content_types.update("application/vnd.openxmlformats-officedocument." + ending for ending in (
        "extended-properties+xml", "custom-properties+xml", "theme+xml", "spreadsheetml.sheet.main+xml",
        "spreadsheetml.worksheet+xml", "spreadsheetml.styles+xml", "spreadsheetml.sharedStrings+xml"))
    for declaration in roots["[Content_Types].xml"]:
        if declaration.get("ContentType") not in accepted_content_types:
            _flag(snapshot, "package_content_type", declaration=dict(declaration.attrib))
    workbook = roots["xl/workbook.xml"]
    styles = Stylesheet.from_tree(roots["xl/styles.xml"])
    formats = {item.numFmtId: item.formatCode for item in styles.numFmts.numFmt}
    strings = list(roots.get("xl/sharedStrings.xml", []))
    relationships = {item.get("Id"): item.get("Target") for item in roots["xl/_rels/workbook.xml.rels"]}
    snapshot["observations"]["workbook_xml"] = parts["xl/workbook.xml"].decode("utf-8-sig")
    snapshot["observations"]["styles_xml"] = parts["xl/styles.xml"].decode("utf-8-sig")
    for defined in workbook.findall(f"{{{NS}}}definedNames/{{{NS}}}definedName"):
        if defined.get("name") not in {"_xlnm._FilterDatabase", "_xlnm.Print_Titles", "_xlnm.Print_Area"}:
            _flag(snapshot, "defined_name", name=defined.get("name"), text=defined.text)
    referenced = set()
    for declaration in workbook.findall(f"{{{NS}}}sheets/{{{NS}}}sheet"):
        target = relationships[declaration.get(f"{{{REL}}}id")]
        part = posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)
        if part in referenced:
            raise ValueError("Repeated worksheet target")
        referenced.add(part)
        root = roots[part]
        sheet = {"name": declaration.get("name"), "visibility": declaration.get("state", "visible"),
                 "cells": {}, "merges": [], "hidden_rows": [], "hidden_columns": [],
                 "observations": {"worksheet_xml": parts[part].decode("utf-8-sig")}}
        dimension = root.find(f"{{{NS}}}dimension")
        if dimension is not None:
            _bounds(dimension.get("ref"))
        sheet["merges"] = [e.get("ref") for e in root.findall(f"{{{NS}}}mergeCells/{{{NS}}}mergeCell")]
        sheet["hidden_rows"] = [int(e.get("r")) for e in root.findall(f"{{{NS}}}sheetData/{{{NS}}}row") if e.get("hidden") in ("1", "true") or float(e.get("ht", "1")) <= 0]
        sheet["hidden_columns"] = [dict(e.attrib) for e in root.findall(f"{{{NS}}}cols/{{{NS}}}col") if e.get("hidden") in ("1", "true") or float(e.get("width", "1")) <= 0]
        sheet_format = root.find(f"{{{NS}}}sheetFormatPr")
        sheet["default_row_hidden"] = sheet_format is not None and (sheet_format.get("zeroHeight") in ("1", "true") or float(sheet_format.get("defaultRowHeight", "1")) <= 0 or float(sheet_format.get("defaultColWidth", "1")) <= 0)
        header = root.find(f"{{{NS}}}headerFooter")
        sheet["observations"]["header_footer"] = "" if header is None else "\n".join(header.itertext())
        if sheet["observations"]["header_footer"].strip():
            _flag(snapshot, "header_footer", sheet=sheet["name"])
        seen = set()
        row_styles = {int(row.get("r")): int(row.get("s", "0")) for row in root.findall(f"{{{NS}}}sheetData/{{{NS}}}row")}
        column_styles = [(int(col.get("min")), int(col.get("max")), int(col.get("style", "0")))
                         for col in root.findall(f"{{{NS}}}cols/{{{NS}}}col")]
        for element in root.findall(f"{{{NS}}}sheetData/{{{NS}}}row/{{{NS}}}c"):
            coordinate = element.get("r")
            _bounds(coordinate)
            if coordinate in seen or len(seen) >= MAX_CELLS:
                raise ValueError("Duplicate or excessive cells")
            seen.add(coordinate)
            style = styles.cellXfs.xf[int(element.get("s", "0"))]
            font = font_snapshot(styles.fonts[style.fontId])
            number_format = formats.get(style.numFmtId, BUILTIN_FORMATS.get(style.numFmtId, "unknown"))
            kind = element.get("t", "n")
            value = element.findtext(f"{{{NS}}}v")
            formula = element.find(f"{{{NS}}}f")
            inline_string = element.find(f"{{{NS}}}is")
            if inline_string is not None and (kind != "inlineStr" or element.find(f"{{{NS}}}v") is not None or formula is not None):
                _flag(snapshot, "cell_representation", sheet=sheet["name"], coordinate=coordinate,
                      detail="inline string conflicts with cell type, value or formula")
            runs = []
            if formula is not None:
                value, kind = "=" + (formula.text or ""), "f"
            elif kind in ("s", "inlineStr"):
                string = strings[int(value)] if kind == "s" else inline_string
                value, runs = _rich_text(string, font)
            elif kind == "n" and value is not None:
                value = float(value) if any(c in value.lower() for c in ".e") else int(value)
                if not math.isfinite(value):
                    raise ValueError("Nonfinite number")
            elif kind == "b":
                if value not in ("0", "1"):
                    _flag(snapshot, "cell_representation", sheet=sheet["name"], coordinate=coordinate,
                          detail="only canonical 0/1 boolean values are supported")
                else:
                    value = value == "1"
            elif kind not in ("n", "e", "str", "d"):
                _flag(snapshot, "cell_type", sheet=sheet["name"], coordinate=coordinate, value_type=kind)
            if value is None or value == "":
                continue
            cell = cell_snapshot(coordinate, value, kind, number_format, font, runs)
            row, column = coordinate_to_tuple(coordinate)
            if element.get("s") is None:
                inherited = {"row_style": row_styles.get(row, 0),
                             "column_styles": [style for left, right, style in column_styles if left <= column <= right and style]}
                if inherited["row_style"] or inherited["column_styles"]:
                    cell["inherited_style"] = inherited
                    _flag(snapshot, "style_inheritance", sheet=sheet["name"], coordinate=coordinate, **inherited)
            cell["source_xml"] = ET.tostring(element, encoding="unicode")
            sheet["cells"][coordinate] = cell
        _finish_sheet(snapshot, sheet)
        snapshot["sheets"].append(sheet)
    if referenced != {name for name in roots if re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml", name)}:
        _flag(snapshot, "unreferenced_worksheet")


def _read_xls(payload, snapshot):
    from xlrd.biffh import biff_count_records
    record_book = xlrd.book.Book()
    record_book.biff2_8_load(file_contents=payload)
    output = io.StringIO()
    try:
        biff_count_records(record_book.mem, record_book.base, record_book.stream_len, output)
    finally:
        record_book.release_resources()
    snapshot["observations"]["biff_records"] = output.getvalue()
    _flag(snapshot, "biff_semantics_unverified", detail="BIFF record counts do not verify formula, extension, object or header/footer semantics")
    book = xlrd.open_workbook(file_contents=payload, formatting_info=True, on_demand=True, ragged_rows=True)
    try:
        for original in book.sheets():
            if original.nrows > MAX_ROWS or original.ncols > MAX_COLUMNS or original.nrows * original.ncols > MAX_CELLS:
                raise ValueError("Worksheet dimensions exceed supported limits")
            sheet = {"name": original.name, "visibility": original.visibility, "cells": {},
                     "merges": [f"{get_column_letter(c1 + 1)}{r1 + 1}:{get_column_letter(c2)}{r2}" for r1, r2, c1, c2 in original.merged_cells],
                     "hidden_rows": [r + 1 for r, info in original.rowinfo_map.items() if info.hidden or info.height == 0],
                     "hidden_columns": [c + 1 for c, info in original.colinfo_map.items() if info.hidden or info.width == 0],
                     "default_row_hidden": original.default_row_hidden,
                     "observations": {"header_footer": "unverified; see BIFF records"}}
            for r in range(original.nrows):
                for c in range(original.row_len(r)):
                    value = original.cell(r, c)
                    if value.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    coordinate = f"{get_column_letter(c + 1)}{r + 1}"
                    xf = book.xf_list[value.xf_index]
                    font = font_snapshot(book.font_list[xf.font_index], biff=True)
                    number_format = book.format_map[xf.format_key].format_str
                    kind = {1: "s", 2: "n", 3: "d", 4: "b", 5: "e"}.get(value.ctype, "unknown")
                    runs = []
                    if value.ctype == xlrd.XL_CELL_TEXT:
                        positions = original.rich_text_runlist_map.get((r, c), [])
                        if positions:
                            if positions[0][0] != 0:
                                positions = [(0, xf.font_index), *positions]
                            for i, (start, index) in enumerate(positions):
                                end = positions[i + 1][0] if i + 1 < len(positions) else len(value.value)
                                runs.append({"text": value.value[start:end], "font": font_snapshot(book.font_list[index], biff=True)})
                    sheet["cells"][coordinate] = cell_snapshot(coordinate, value.value, kind, number_format, font, runs)
            _finish_sheet(snapshot, sheet)
            snapshot["sheets"].append(sheet)
    finally:
        book.release_resources()


def read_workbook(payload, suffix):
    """Return original cell content and audit observations; perform no I/O."""
    if not isinstance(payload, bytes) or not payload or len(payload) > MAX_INPUT:
        raise ValueError("Workbook must be nonempty bytes at most 10 MiB")
    if suffix not in (".xls", ".xlsx"):
        raise ValueError("Only .xls and .xlsx are supported")
    if (suffix == ".xls" and not payload.startswith(bytes.fromhex("D0CF11E0A1B11AE1"))) or (suffix == ".xlsx" and not payload.startswith(b"PK\x03\x04")):
        raise ValueError("Workbook signature does not match suffix")
    snapshot = {"sha256": hashlib.sha256(payload).hexdigest(), "format": suffix,
                "reader": {"openpyxl": openpyxl.__version__, "xlrd": xlrd.__version__},
                "sheets": [], "unsupported": [], "observations": {}}
    try:
        (_read_xlsx if suffix == ".xlsx" else _read_xls)(payload, snapshot)
    except (ValueError, KeyError, IndexError, TypeError, AttributeError, ET.ParseError, zipfile.BadZipFile, xlrd.XLRDError) as error:
        raise ValueError(f"Cannot safely read workbook: {error}") from error
    if not snapshot["sheets"]:
        raise ValueError("Workbook contains no worksheets")
    return snapshot


def project_workbook(snapshot):
    """Project six-column row context, preserving original text and merge origin."""
    if snapshot["unsupported"]:
        raise ValueError("Workbook has unsupported structures; projection refused")
    text, locations, offset = [], [], 0
    for sheet_index, sheet in enumerate(snapshot["sheets"]):
        ranges = [(reference, _bounds(reference)) for reference in sheet["merges"]]
        if not sheet["cells"]:
            block = f"[工作表 {sheet['name']} | 无非空单元格]\n[合并范围] {', '.join(sheet['merges']) or '无'}\n\n"
            text.append(block)
            locations.append({"id": f"{snapshot['sha256']}:{sheet_index}:empty", "sheet": sheet["name"],
                              "row": None, "start": offset, "end": offset + len(block),
                              "kind": "empty_sheet", "cells": [], "merges": sheet["merges"][:]})
            offset += len(block)
            continue
        last_row = max(coordinate_to_tuple(c)[0] for c in sheet["cells"])
        last_row = max(last_row, *(bounds[3] for _, bounds in ranges)) if ranges else last_row
        for row in range(1, last_row + 1):
            entries, lines = [], [f"[工作表 {sheet['name']} | 行 {row}]"]
            for col in range(1, 7):
                coordinate = f"{get_column_letter(col)}{row}"
                source, anchor = coordinate, None
                for _, (left, top, right, bottom) in ranges:
                    if left <= col <= right and top <= row <= bottom:
                        source = anchor = f"{get_column_letter(left)}{top}"
                        break
                cell = sheet["cells"].get(source, {})
                entry = {"coordinate": coordinate, "source_coordinate": source, "merge_anchor": anchor,
                         "raw_text": cell.get("raw_text", ""), "display_text": cell.get("display_text", "")}
                entries.append(entry)
                origin = f"；合并锚点 {source}" if anchor else ""
                label = f" {HEADERS[col - 1]}" if row > 3 else ""
                lines.append(f"[{coordinate}{label}{origin}] {entry['display_text']}")
            full_width = next((reference for reference, (left, top, right, bottom) in ranges
                               if left == 1 and right == 6 and top <= row <= bottom), None)
            if full_width:
                lines = [lines[0], f"[整行合并 {full_width}；锚点 {entries[0]['source_coordinate']}] {entries[0]['display_text']}"]
            block = "\n".join(lines) + "\n\n"
            locations.append({"id": f"{snapshot['sha256']}:{sheet_index}:{row}", "sheet": sheet["name"],
                              "row": row, "start": offset, "end": offset + len(block), "cells": entries,
                              "kind": "merged_row" if full_width else "row", "headers": list(HEADERS) if row > 3 and not full_width else []})
            text.append(block)
            offset += len(block)
    return "".join(text), locations
